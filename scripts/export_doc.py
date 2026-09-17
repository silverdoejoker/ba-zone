"""
Export Markdown with Mermaid diagrams to HTML & DOCX (Word)
Author: BA Zone Toolkit
"""
import os
import re
import sys
import base64
import requests
import markdown
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def render_mermaid_to_png(mermaid_code, output_png_path):
    """Render Mermaid code to PNG using Kroki or Mermaid.ink"""
    # Clean up mermaid code
    mermaid_code = mermaid_code.strip()
    
    # Method 1: Kroki
    try:
        url = "https://kroki.io/mermaid/png"
        resp = requests.post(
            url,
            data=mermaid_code.encode("utf-8"),
            headers={"Content-Type": "text/plain; charset=utf-8"},
            timeout=15
        )
        if resp.status_code == 200 and len(resp.content) > 100:
            with open(output_png_path, "wb") as f:
                f.write(resp.content)
            return True
    except Exception as e:
        print(f"Kroki error: {e}")

    # Method 2: Mermaid.ink fallback
    try:
        encoded = base64.urlsafe_b64encode(mermaid_code.encode("utf-8")).decode("ascii")
        url = f"https://mermaid.ink/img/{encoded}"
        resp = requests.get(url, timeout=15)
        if resp.status_code == 200 and len(resp.content) > 100:
            with open(output_png_path, "wb") as f:
                f.write(resp.content)
            return True
    except Exception as e:
        print(f"Mermaid.ink error: {e}")

    return False


def export_to_html(md_path, html_path):
    """Export Markdown to standalone HTML with interactive Mermaid.js rendering"""
    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    # Replace ```mermaid with <div class="mermaid">
    def replace_mermaid(match):
        code = match.group(1).strip()
        return f'<div class="mermaid">\n{code}\n</div>'

    processed_md = re.sub(r'```mermaid\s*\n(.*?)\n```', replace_mermaid, content, flags=re.DOTALL)

    # Convert markdown to html
    html_body = markdown.markdown(
        processed_md,
        extensions=['extra', 'tables', 'toc', 'nl2br']
    )

    title_match = re.search(r'^#\s+(.+)$', content, re.MULTILINE)
    title = title_match.group(1) if title_match else "Business Requirements Document"

    html_template = f"""<!DOCTYPE html>
<html lang="vi">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{title}</title>
    <!-- Mermaid JS -->
    <script src="https://cdn.jsdelivr.net/npm/mermaid@10/dist/mermaid.min.js"></script>
    <script>
        mermaid.initialize({{
            startOnLoad: true,
            theme: 'default',
            securityLevel: 'loose',
            flowchart: {{ useMaxWidth: true, htmlLabels: true, curve: 'basis' }},
            sequence: {{ useMaxWidth: true, showSequenceNumbers: true }}
        }});
    </script>
    <style>
        @import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&display=swap');
        
        :root {{
            --primary: #0052cc;
            --primary-dark: #0747a6;
            --text-main: #172b4d;
            --text-sub: #5e6c84;
            --bg-page: #f4f5f7;
            --bg-card: #ffffff;
            --border: #ebecf0;
        }}

        body {{
            font-family: 'Inter', -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
            background-color: var(--bg-page);
            color: var(--text-main);
            line-height: 1.6;
            margin: 0;
            padding: 20px;
        }}

        .container {{
            max-width: 960px;
            margin: 0 auto;
            background: var(--bg-card);
            padding: 50px 60px;
            border-radius: 8px;
            box-shadow: 0 4px 12px rgba(9, 30, 66, 0.08);
        }}

        /* Print Action Bar */
        .action-bar {{
            position: fixed;
            top: 20px;
            right: 20px;
            z-index: 9999;
            background: white;
            padding: 8px 12px;
            border-radius: 30px;
            box-shadow: 0 4px 16px rgba(0,0,0,0.15);
            display: flex;
            gap: 10px;
        }}
        .btn-print {{
            background: #0052cc;
            color: white;
            border: none;
            padding: 8px 16px;
            border-radius: 20px;
            font-size: 13px;
            font-weight: 600;
            cursor: pointer;
            display: flex;
            align-items: center;
            gap: 6px;
            transition: 0.2s;
        }}
        .btn-print:hover {{
            background: #0747a6;
            transform: translateY(-1px);
        }}

        h1 {{
            color: #0c2340;
            font-size: 26px;
            border-bottom: 2px solid #0052cc;
            padding-bottom: 12px;
            margin-top: 20px;
            margin-bottom: 8px;
        }}
        h2 {{
            color: #172b4d;
            font-size: 20px;
            margin-top: 36px;
            margin-bottom: 12px;
            border-left: 4px solid #0052cc;
            padding-left: 10px;
        }}
        h3 {{
            color: #344563;
            font-size: 16px;
            margin-top: 24px;
            margin-bottom: 8px;
        }}

        blockquote {{
            margin: 16px 0;
            padding: 12px 18px;
            background: #f4f5f7;
            border-left: 4px solid #4c9aff;
            color: #42526e;
            border-radius: 0 4px 4px 0;
            font-size: 13.5px;
        }}

        table {{
            width: 100%;
            border-collapse: collapse;
            margin: 20px 0;
            font-size: 13px;
        }}
        th, td {{
            border: 1px solid #dfe1e6;
            padding: 10px 12px;
            text-align: left;
            vertical-align: top;
        }}
        th {{
            background-color: #f4f5f7;
            color: #172b4d;
            font-weight: 600;
        }}
        tr:nth-child(even) {{
            background-color: #fafbfc;
        }}

        .mermaid {{
            background: #ffffff;
            border: 1px solid #ebecf0;
            border-radius: 8px;
            padding: 20px;
            margin: 24px 0;
            display: flex;
            justify-content: center;
            box-shadow: inset 0 0 4px rgba(0,0,0,0.02);
            overflow-x: auto;
        }}

        code {{
            background: #f4f5f7;
            color: #bf2600;
            padding: 2px 6px;
            border-radius: 3px;
            font-size: 12.5px;
            font-family: Consolas, monospace;
        }}

        hr {{
            border: none;
            border-top: 1px solid #ebecf0;
            margin: 30px 0;
        }}

        /* Print Media Styles */
        @media print {{
            body {{
                background: white;
                padding: 0;
                color: black;
            }}
            .container {{
                box-shadow: none;
                padding: 0;
                max-width: 100%;
            }}
            .action-bar {{
                display: none;
            }}
            @page {{
                size: A4 portrait;
                margin: 15mm 15mm 15mm 15mm;
            }}
            h1, h2, h3 {{
                page-break-after: avoid;
            }}
            table, .mermaid {{
                page-break-inside: avoid;
            }}
        }}
    </style>
</head>
<body>
    <div class="action-bar">
        <button class="btn-print" onclick="window.print()">🖨️ In / Lưu PDF (Ctrl + P)</button>
    </div>
    <div class="container">
        {html_body}
    </div>
</body>
</html>
"""
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_template)
    print(f"Exported HTML to: {html_path}")


def set_cell_background(cell, fill_color):
    """Set cell background color (hex string e.g. 'EBF1F5')"""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner cell padding"""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def export_to_docx(md_path, docx_path):
    """Export Markdown with embedded Mermaid PNG diagrams to Word (.docx)"""
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()

    # Set page margins (1 inch = 2.54 cm)
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    # Styles
    navy = RGBColor(12, 35, 64)
    blue = RGBColor(0, 82, 204)
    dark_gray = RGBColor(50, 50, 50)

    # State machine for parsing markdown lines
    in_mermaid = False
    mermaid_buffer = []
    mermaid_count = 0
    in_table = False
    table_rows = []

    temp_img_dir = os.path.join(os.path.dirname(docx_path), "temp_diagrams")
    os.makedirs(temp_img_dir, exist_ok=True)

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        
        # Determine cols
        num_cols = max(len(r) for r in table_rows)
        tbl = doc.add_table(rows=len(table_rows), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        for r_idx, row in enumerate(table_rows):
            for c_idx in range(num_cols):
                cell = tbl.cell(r_idx, c_idx)
                text = row[c_idx] if c_idx < len(row) else ""
                cell.text = text.strip()
                set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                
                # Header row styling
                if r_idx == 0:
                    set_cell_background(cell, "0052CC")
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                            run.font.size = Pt(9.5)
                else:
                    if r_idx % 2 == 0:
                        set_cell_background(cell, "F4F5F7")
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(9)
                            run.font.color.rgb = dark_gray

        doc.add_paragraph() # Spacing
        table_rows = []
        in_table = False

    for line in lines:
        raw = line.rstrip("\r\n")

        # Check mermaid block
        if raw.strip().startswith("```mermaid"):
            if in_table:
                flush_table()
            in_mermaid = True
            mermaid_buffer = []
            continue

        if in_mermaid:
            if raw.strip() == "```":
                in_mermaid = False
                mermaid_count += 1
                code = "\n".join(mermaid_buffer)
                img_path = os.path.join(temp_img_dir, f"diagram_{mermaid_count}.png")
                
                print(f"Rendering Mermaid Diagram #{mermaid_count}...")
                success = render_mermaid_to_png(code, img_path)
                if success and os.path.exists(img_path):
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p_img.add_run()
                    run.add_picture(img_path, width=Inches(6.2))
                    
                    # Caption
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r_cap = p_cap.add_run(f"Sơ đồ {mermaid_count}: Biểu đồ luồng tích hợp / quy trình hệ thống")
                    r_cap.font.italic = True
                    r_cap.font.size = Pt(8.5)
                    r_cap.font.color.rgb = RGBColor(120, 120, 120)
                else:
                    # Fallback text
                    p_fb = doc.add_paragraph()
                    p_fb.add_run(f"[Diagram #{mermaid_count}: Xem sơ đồ trực tiếp trên bản HTML / Markdown]").font.italic = True
                
                doc.add_paragraph()
            else:
                mermaid_buffer.append(raw)
            continue

        # Check tables
        if "|" in raw and "-|-" in raw:
            # Separator line, ignore
            continue
        elif raw.strip().startswith("|") and raw.strip().endswith("|"):
            in_table = True
            cols = [c.strip() for c in raw.strip().split("|")[1:-1]]
            table_rows.append(cols)
            continue
        else:
            if in_table:
                flush_table()

        # Check headings
        if raw.startswith("# "):
            p = doc.add_paragraph()
            run = p.add_run(raw[2:].strip())
            run.font.bold = True
            run.font.size = Pt(18)
            run.font.color.rgb = navy
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
        elif raw.startswith("## "):
            p = doc.add_paragraph()
            run = p.add_run(raw[3:].strip())
            run.font.bold = True
            run.font.size = Pt(14)
            run.font.color.rgb = blue
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
        elif raw.startswith("### "):
            p = doc.add_paragraph()
            run = p.add_run(raw[4:].strip())
            run.font.bold = True
            run.font.size = Pt(11.5)
            run.font.color.rgb = navy
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        elif raw.startswith("> "):
            # Blockquote
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            run = p.add_run(raw[2:].strip())
            run.font.italic = True
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(90, 100, 115)
        elif raw.startswith("* ") or raw.startswith("- "):
            p = doc.add_paragraph(style='List Bullet')
            # Clean bold text inside list item
            text = raw[2:].strip()
            add_formatted_text(p, text)
        elif re.match(r'^\d+\.\s', raw):
            p = doc.add_paragraph(style='List Number')
            text = re.sub(r'^\d+\.\s', '', raw).strip()
            add_formatted_text(p, text)
        elif raw.strip() == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
        elif raw.strip() == "":
            pass
        else:
            p = doc.add_paragraph()
            add_formatted_text(p, raw.strip())
            p.paragraph_format.space_after = Pt(4)

    if in_table:
        flush_table()

    doc.save(docx_path)
    print(f"Exported DOCX to: {docx_path}")


def add_formatted_text(paragraph, text):
    """Parse inline **bold** and *italic* into docx runs"""
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            r.font.bold = True
            r.font.size = Pt(10)
        elif part.startswith("*") and part.endswith("*"):
            r = paragraph.add_run(part[1:-1])
            r.font.italic = True
            r.font.size = Pt(10)
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(170, 30, 30)
        else:
            r = paragraph.add_run(part)
            r.font.size = Pt(10)


if __name__ == "__main__":
    target_md = sys.argv[1] if len(sys.argv) > 1 else r"d:\repo\ba-zone\docs\outputs\training-attendance-system\BRD_TRAINING_ATTENDANCE_SYSTEM.md"
    base_name = os.path.splitext(target_md)[0]
    html_out = base_name + ".html"
    docx_out = base_name + ".docx"

    print(f"--- EXPORTING: {target_md} ---")
    export_to_html(target_md, html_out)
    export_to_docx(target_md, docx_out)
    print("ALL EXPORTS COMPLETED SUCCESSFULLY!")
